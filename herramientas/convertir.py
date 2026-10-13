"""
Convierte la "Tabla dilución medicamentos" (Excel) en datos/medicamentos.json,
que es lo que lee la página.

Uso (desde la raíz del proyecto):
    py herramientas/convertir.py
    py herramientas/convertir.py "ruta/a/otro excel.xlsx"

Cada hoja de medicamento tiene pares etiqueta -> valor (etiqueta en una columna,
valor en la siguiente). Si una hoja tiene varias marcas lado a lado, cada bloque
de etiquetas se vuelve una "presentación" de la misma ficha.
"""

import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
EXCEL_POR_DEFECTO = RAIZ / "datos" / "fuente" / "Tabla dilución medicamentos 2024.xlsx"
SALIDA = RAIZ / "datos" / "medicamentos.json"

# --CAMPOS DE LA FICHA, EN EL ORDEN EN QUE SE MUESTRAN
# clave -> (titulo para la web, patron de la etiqueta del Excel ya compactada: sin tildes, espacios ni simbolos)
CAMPOS = {
    "reconstitucion": ("Volumen de reconstitución", r"volumenreconstitucion"),
    "estab_recon_amb": ("Estabilidad del vial reconstituido (T amb ≤ 30 °C)", r"estabilidadampollareconstituida(tamb.*)?"),
    "estab_recon_frio": ("Estabilidad del vial reconstituido (2-8 °C)", r"estabilidadampollareconstituida(t)?28c?"),
    "vol_normal": ("Volumen normal de dilución", r"volumennormaldedilucion"),
    "conc_estandar": ("Concentración de dilución estándar", r"concentraciondedilucionestandar"),
    "vol_restringido": ("Volumen restringido de dilución", r"volumenrestringidodedilucion"),
    "conc_restringido": ("Concentración de volumen restringido", r"concentraciondevolumenrestringido"),
    "estab_sol_amb": ("Estabilidad en solución (T amb ≤ 30 °C)", r"estabilidaddelmedicamento(ensolucion)?(tamb.*)?"),
    "estab_sol_frio": ("Estabilidad en solución (2-8 °C)", r"estabilidaddelmedicamento(ensolucion)?(t)?28c?"),
    "soluciones_compatibles": ("Soluciones compatibles para dilución", r"solucionescompatiblesparadilucion"),
    "tiempo_infusion": ("Tiempo de infusión", r"tiempodeinfusion"),
    "dosis_mantenimiento": ("Dosis de mantenimiento en infusión continua", r"dosisdemantenimientoeninfusioncontinua"),
    "bolos": ("Necesidad de bolos", r"necesidaddebolos"),
    "dosis_maxima": ("Dosis máxima permitida", r"dosismaximapermitida(pordia)?"),
    "osmolaridad": ("Osmolaridad en solución", r"osmolaridad(ensolucion)?"),
    "ph": ("pH en solución", r"ph(ensolucion)?"),
    "precauciones": ("Precauciones de administración", r"precaucionesdeadministracion"),
    "no_compatibles": ("Medicamentos NO compatibles para mezcla o infusión en Y", r"medicamentosnocompatiblesparamezclaoinfusioneny"),
    "efectos": ("Efectos secundarios relevantes", r"(efectossecundarios|reaccionesadversas)relevantes"),
    "observaciones": ("Observaciones", r"observaciones"),
}

# --NOMBRES DE HOJA ABREVIADOS -> NOMBRE PARA MOSTRAR (TOMADOS DEL INDICE DEL EXCEL)
NOMBRES = {
    "imipinemcilas500mg": "Imipenem/Cilastatina 500 mg",
    "imipinemcilas1g": "Imipenem/Cilastatina 1 g",
    "trimetoprimsulfa80400mg": "Trimetoprim/Sulfametoxazol 80/400 mg",
    "fenitoina250mg5ml": "Fenitoína 250 mg/5 mL",
    "colistimetato1mui": "Colistimetato 1.000.000 UI",
    "colistimetato2mui": "Colistimetato 2.000.000 UI",
    "octreotride01g": "Octreotide 0,1 mg/1 mL",
    "remifentanyl2mg": "Remifentanilo 2 mg",
    "tiopenthal1g": "Tiopental 1 g",
    "actylise50mg": "Actilyse 50 mg",
    "lidocaina10ml": "Lidocaína (clorhidrato) 10 mL",
    "ceftazidimaavibactam": "Ceftazidima 2 g + Avibactam 0,5 g",
    "penicilinacristalina5000000": "Penicilina Cristalina 5.000.000 UI",
    "fentanil500mcg": "Fentanilo 500 mcg",
    "meropenem1g": "Meropenem 1 g",
    "anfotericinaliposomal": "Anfotericina B liposomal 50 mg",
}

# --HOJAS QUE NO SON FICHAS DE MEDICAMENTO
HOJAS_ESPECIALES = {"portada", "indice", "soluciones", "insulina suplementaria"}

# --ELECTROLITOS (SEGUN EL INDICE DEL EXCEL), POR PREFIJO DEL NOMBRE DE LA HOJA
ELECTROLITOS = ["bicarbonato", "calcio gluconato", "fosfato de potasio", "magnesio sulfato", "sodio cloruro", "potasio"]


def normalizar(texto):
    """Minúsculas, sin tildes y con espacios simples: para comparar etiquetas y nombres."""
    texto = unicodedata.normalize("NFD", str(texto).lower())
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", texto).strip()


def compacto(texto):
    """Solo letras y números: para emparejar nombres escritos distinto (Piperacilina+Tazobactam / Piperacilina Tazobactam)."""
    return re.sub(r"[^a-z0-9]", "", normalizar(texto))


def limpiar(valor):
    """Texto del Excel listo para mostrar: conserva saltos de línea, quita espacios sobrantes."""
    if valor is None:
        return ""
    if isinstance(valor, float):
        valor = int(valor) if valor.is_integer() else round(valor, 2)
    lineas = [re.sub(r"[ \t]+", " ", l).strip() for l in str(valor).replace("\r", "").split("\n")]
    texto = "\n".join(l for l in lineas if l)
    return texto.strip(" ,")


def slug(texto):
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", normalizar(texto))).strip("-")


def campo_de_etiqueta(texto):
    """Clave del campo si el texto es una etiqueta de la ficha; None si es un valor u otra cosa."""
    k = compacto(texto)
    if not k or len(k) > 70:
        return None
    for clave, (_, patron) in CAMPOS.items():
        if re.fullmatch(patron, k):
            return clave
    return None


def es_texto_ignorable(texto):
    n = normalizar(texto)
    return n.startswith("tabla de dilucion") or n.startswith("volver al indice") or n in ("false", "true")


def leer_ficha(ws):
    """Devuelve la ficha de una hoja de medicamento, o None si la hoja no tiene etiquetas."""
    etiquetas = []  # (fila, columna, campo)
    for fila in ws.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str):
                campo = campo_de_etiqueta(celda.value)
                if campo:
                    etiquetas.append((celda.row, celda.column, campo))
    if not etiquetas:
        return None

    usadas = set()
    columnas = sorted({c for _, c, _ in etiquetas})
    presentaciones = []
    for col in columnas:
        del_bloque = [(f, campo) for f, c, campo in etiquetas if c == col]
        primera_fila = min(f for f, _ in del_bloque)

        # --TITULO: PRIMER TEXTO EN LA COLUMNA DE ETIQUETAS POR ENCIMA DEL BLOQUE
        titulo = ""
        for f in range(primera_fila - 1, 0, -1):
            v = ws.cell(f, col).value
            if v is not None and limpiar(v) and not es_texto_ignorable(v):
                titulo = limpiar(v).replace("\n", " ")
                usadas.add((f, col))
                break

        # --SI A ESTA MARCA LE FALTA UNA ETIQUETA QUE OTRA MARCA SI TIENE EN LA MISMA FILA, SE HEREDA
        filas_bloque = {f for f, _ in del_bloque}
        for f, c, campo in etiquetas:
            if c != col and f not in filas_bloque and f > primera_fila                     and not limpiar(ws.cell(f, col).value) and limpiar(ws.cell(f, col + 1).value):
                del_bloque.append((f, campo))
                filas_bloque.add(f)

        datos = {}
        for f, campo in del_bloque:
            usadas.add((f, col))
            usadas.add((f, col + 1))
            valor = limpiar(ws.cell(f, col + 1).value)
            if valor:
                datos[campo] = valor
        presentaciones.append({"titulo": titulo, "datos": datos})

    # --NOTAS: TEXTOS SUELTOS DE LA HOJA (FICHA TECNICA, ENLACES, AJUSTES CON FECHA)
    notas = []
    valores = {v for p in presentaciones for v in p["datos"].values()}
    for fila in ws.iter_rows():
        for celda in fila:
            if (celda.row, celda.column) in usadas or celda.value is None:
                continue
            texto = limpiar(celda.value)
            if len(texto) > 2 and texto not in valores and not es_texto_ignorable(texto) and not isinstance(celda.value, (int, float)):
                if texto not in notas:
                    notas.append(texto)

    return {"presentaciones": presentaciones, "notas": notas}


def nombre_marca(titulo):
    """Saca la marca entre paréntesis del título, p. ej. '(ROCHE) CEFTRIAXONA 1 g' -> 'ROCHE'."""
    m = re.findall(r"\(([A-Za-zÁÉÍÓÚÑ][A-Za-zÁÉÍÓÚÑáéíóúñ .&-]{2,})\)", titulo)
    return m[-1].strip() if m else ""


def forma(titulos):
    t = normalizar(" ".join(titulos))
    if "liofilizado" in t or "polvo" in t:
        return "Polvo liofilizado"
    if "inyectable" in t or "solucion" in t or "/" in t:
        return "Solución inyectable"
    return ""


def leer_periferica(ws):
    normas, medicamentos, nota = [], [], ""
    for fila in ws.iter_rows(min_row=5):
        for celda in fila:
            if celda.value is None or isinstance(celda.value, (int, float, bool)):
                continue
            texto = limpiar(celda.value)
            if celda.column_letter == "C":
                normas.append(texto)
            elif celda.column_letter == "F":
                medicamentos.append(texto)
            elif celda.column_letter == "B" and len(texto) > 40:
                nota = texto
    return {"titulo": limpiar(ws["A2"].value), "normas": normas, "medicamentos": medicamentos, "nota": nota}


def leer_subcutanea(ws):
    normas, aptos, tabla = [], [], []
    for f in range(6, ws.max_row + 1):
        c = limpiar(ws.cell(f, 3).value)
        if c and f <= 8:
            normas.append(c)
        elif c:
            aptos.append({
                "nombre": c,
                "nota": limpiar(ws.cell(f, 4).value),
                "enlace": limpiar(ws.cell(f, 6).value),
            })
        farmaco = limpiar(ws.cell(f, 10).value)
        if farmaco:
            tabla.append({
                "farmaco": farmaco,
                "indicacion": limpiar(ws.cell(f, 11).value),
                "posologia": limpiar(ws.cell(f, 12).value),
            })
    return {"titulo": limpiar(ws["A2"].value), "normas": normas, "aptos": aptos, "tabla": tabla}


def leer_soluciones(ws):
    """Hoja SOLUCIONES: soluciones salinas, solución polarizante, mezcla de potasio 1:1 y soluciones base."""
    salinas = []
    for f in range(5, 9):
        salinas.append({
            "solucion": limpiar(ws.cell(f, 2).value),
            "dilucion": limpiar(ws.cell(f, 3).value),
            "osmolaridad": limpiar(ws.cell(f, 5).value),
        })
    precaucion_salinas = limpiar(ws["F5"].value)

    polarizante = []
    for f in range(11, 19):
        etq, val = limpiar(ws.cell(f, 3).value), limpiar(ws.cell(f, 5).value)
        if etq:
            polarizante.append({"campo": etq, "valor": val})

    encabezados = [limpiar(ws.cell(6, c).value) for c in range(9, 15)]
    potasio = []
    for f in range(7, 17):
        potasio.append([limpiar(ws.cell(f, c).value) for c in range(9, 15)])
    nota_potasio = limpiar(ws["I17"].value)

    base = []
    for f in range(20, ws.max_row + 1):
        nombre = limpiar(ws.cell(f, 2).value)
        if nombre and normalizar(nombre) not in ("osmolaridad", "ph"):
            base.append({
                "solucion": nombre,
                "osmolaridad": limpiar(ws.cell(f + 2, 5).value),
                "ph": limpiar(ws.cell(f + 3, 5).value),
            })
    return {
        "salinas": salinas, "precaucion_salinas": precaucion_salinas,
        "polarizante": polarizante,
        "potasio": {"encabezados": encabezados, "filas": potasio, "nota": nota_potasio},
        "base": base,
    }


def leer_insulina(ws):
    encabezados = [limpiar(ws.cell(4, c).value).replace("\n", " ") for c in range(2, 6)]
    filas = []
    for f in range(5, 12):
        filas.append([limpiar(ws.cell(f, c).value) for c in range(2, 6)])
    return {"titulo": limpiar(ws["B2"].value), "encabezados": encabezados, "filas": filas, "nota": limpiar(ws["B12"].value)}


def leer_portada(ws):
    textos = [limpiar(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1)]
    textos = [t for t in textos if t]
    elaborado = [t for t in textos if normalizar(t).startswith("elaborado") or normalizar(t).startswith("q.f")]
    aprobado = next((t for t in textos if normalizar(t).startswith("aprobado")), "")
    referencias = [t for t in textos if t.startswith("http")]
    autores = []
    for t in elaborado:
        for linea in t.split("\n"):
            linea = re.sub(r"^elaborado por", "", linea, flags=re.I).strip()
            if linea:
                autores.append(linea)
    return {
        "titulo": textos[0] if textos else "",
        "autores": autores,
        "aprobado": aprobado.replace("\n", " · ").replace("Aprobado por ·", "Aprobado por"),
        "referencias": referencias,
    }


def leer_indice(ws):
    """Entradas del índice: [{nombre, antibiotico, forma}]. La forma sale del encabezado de su columna."""
    entradas, encabezados, antibioticos = [], {}, True
    for fila in ws.iter_rows():
        for celda in fila:
            if not isinstance(celda.value, str) or not celda.value.strip():
                continue
            texto, n = limpiar(celda.value), normalizar(celda.value)
            if n.startswith("tabla de dilucion"):
                antibioticos = "antibiotico" in n
            elif n in ("polvo liofilizado", "solucion inyectable", "electrolitos", "soluciones"):
                encabezados[celda.column] = texto.capitalize()
            else:
                entradas.append({"nombre": texto, "antibiotico": antibioticos, "forma": encabezados.get(celda.column, "")})
    return entradas


def primera_palabra(texto):
    palabras = re.findall(r"[a-z]+", normalizar(texto))
    palabras = [p for p in palabras if p not in ("clorhidrato", "de")]
    return palabras[0] if palabras else ""


def buscar_en_indice(indice, nombre):
    a = primera_palabra(nombre)
    for e in indice:
        b = primera_palabra(e["nombre"])
        if len(a) >= 5 and len(b) >= 5 and (a.startswith(b[:6]) or b.startswith(a[:6])):
            return e
    return None


def coincide(nombre_lista, nombre_ficha):
    """True si un nombre de las listas (periférica, subcutánea) corresponde a la ficha."""
    a, b = compacto(nombre_lista), compacto(nombre_ficha)
    if len(a) < 5:
        return False
    return b.startswith(a) or b.startswith(a[:-1]) and len(a) >= 7


def main():
    ruta_excel = Path(sys.argv[1]) if len(sys.argv) > 1 else EXCEL_POR_DEFECTO
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)

    medicamentos, especiales, ignoradas = [], {}, []
    indice = next((leer_indice(ws) for ws in wb.worksheets if normalizar(ws.title) == "indice"), [])
    for ws in wb.worksheets:
        n = normalizar(ws.title)
        if n == "portada":
            especiales["fuente"] = leer_portada(ws)
        elif n == "soluciones":
            especiales["soluciones"] = leer_soluciones(ws)
        elif n == "insulina suplementaria":
            especiales["insulina"] = leer_insulina(ws)
        elif n.startswith("listado via periferica"):
            especiales["periferica"] = leer_periferica(ws)
        elif n.startswith("listado via subcutanea"):
            especiales["subcutanea"] = leer_subcutanea(ws)
        elif n in HOJAS_ESPECIALES:
            continue
        else:
            ficha = leer_ficha(ws)
            if not ficha:
                ignoradas.append(ws.title)
                continue
            encabezado = normalizar(" ".join(limpiar(ws.cell(1, c).value) for c in range(1, 6)))
            titulos = [p["titulo"] for p in ficha["presentaciones"]]
            nombre = NOMBRES.get(compacto(ws.title), ws.title.strip())
            en_indice = buscar_en_indice(indice, nombre)
            categoria = "Antibiótico" if "antibiotico" in encabezado or (en_indice and en_indice["antibiotico"]) else "Otro"
            if any(n.startswith(e) for e in ELECTROLITOS):
                categoria = "Electrolito"
            forma_ficha = forma(titulos)
            if en_indice and en_indice["forma"] in ("Polvo liofilizado", "Solución inyectable"):
                forma_ficha = en_indice["forma"]
            if not forma_ficha:
                forma_ficha = "Solución inyectable"
            for p in ficha["presentaciones"]:
                p["marca"] = nombre_marca(p["titulo"])
            medicamentos.append({
                "id": slug(nombre),
                "nombre": nombre,
                "categoria": categoria,
                "forma": forma_ficha,
                "presentaciones": ficha["presentaciones"],
                "notas": ficha["notas"],
                "hoja": ws.title,
            })

    # --MARCAS DE VIA PERIFERICA CON PRECAUCION Y USO SUBCUTANEO
    periferica = especiales.get("periferica", {}).get("medicamentos", [])
    subcutanea = [t["farmaco"] for t in especiales.get("subcutanea", {}).get("tabla", [])]
    subcutanea += [a["nombre"] for a in especiales.get("subcutanea", {}).get("aptos", [])]
    for m in medicamentos:
        m["periferica_precaucion"] = any(coincide(p, m["nombre"]) for p in periferica)
        m["subcutanea"] = any(coincide(s, m["nombre"]) for s in subcutanea)

    medicamentos.sort(key=lambda m: normalizar(m["nombre"]))

    # --IDS UNICOS
    vistos = {}
    for m in medicamentos:
        if m["id"] in vistos:
            vistos[m["id"]] += 1
            m["id"] = f'{m["id"]}-{vistos[m["id"]]}'
        else:
            vistos[m["id"]] = 1

    salida = {
        "generado": date.today().isoformat(),
        "archivo_fuente": ruta_excel.name,
        "campos": [{"clave": k, "titulo": t} for k, (t, _) in CAMPOS.items()],
        "medicamentos": medicamentos,
        **especiales,
    }
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    SALIDA.write_text(json.dumps(salida, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"{len(medicamentos)} medicamentos -> {SALIDA.relative_to(RAIZ)}")
    print(f"Hojas vacías ignoradas: {', '.join(ignoradas) or 'ninguna'}")
    sin_titulo = [m["nombre"] for m in medicamentos for p in m["presentaciones"] if not p["titulo"]]
    if sin_titulo:
        print("Presentaciones sin título:", ", ".join(sin_titulo))


if __name__ == "__main__":
    main()
